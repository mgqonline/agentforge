import pandas as pd
import subprocess
import sys
import os

def install_and_import(package, import_name):
    try:
        __import__(import_name)
    except ImportError:
        print(f"Installing {package}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package])

# 确保必要的库存在
install_and_import("fpdf", "fpdf")
install_and_import("python-docx", "docx")
install_and_import("openpyxl", "openpyxl")

from fpdf import FPDF
import docx

# ==========================================
# 1. 生成 PDF: 财务报告
# ==========================================
print("正在生成 Q3_Financial_Report.pdf ...")
pdf = FPDF()
pdf.add_page()
# 为了避免不同系统缺少中文字体报错，我们使用纯英文生成演示文档
pdf.set_font("Arial", size=16, style='B')
pdf.cell(200, 10, txt="Q3 Financial Report 2026", ln=1, align='C')
pdf.set_font("Arial", size=12)
pdf.cell(200, 10, txt="--------------------------------------------------", ln=1, align='C')
pdf.cell(200, 10, txt="Total Revenue: $5,200,000", ln=1)
pdf.cell(200, 10, txt="Net Profit: $1,450,000", ln=1)
pdf.cell(200, 10, txt="Major expenditures include R&D and Marketing.", ln=1)
pdf.cell(200, 10, txt="", ln=1)
pdf.cell(200, 10, txt="Summary: The company experienced a 15% growth in Q3.", ln=1)
pdf.output("Q3_Financial_Report.pdf")

# ==========================================
# 2. 生成 Word: 劳动合同
# ==========================================
print("正在生成 Labor_Contract.docx ...")
doc = docx.Document()
doc.add_heading('Standard Labor Contract', 0)
doc.add_paragraph('This employment agreement is made and effective as of today.')
doc.add_heading('1. Position and Duties', level=1)
doc.add_paragraph('The Employee will serve as a Senior AI Engineer. Responsibilities include building RAG pipelines and optimizing prompts.')
doc.add_heading('2. Compensation', level=1)
doc.add_paragraph('The Employee will receive a base salary of $150,000 per year, paid monthly.')
doc.save("Labor_Contract.docx")

# ==========================================
# 3. 生成 Excel: 销售数据
# ==========================================
print("正在生成 Sales_Data_2026.xlsx ...")
data = {
    'Employee ID': ['E001', 'E002', 'E003', 'E004'],
    'Name': ['Alice Smith', 'Bob Jones', 'Charlie Brown', 'Diana Prince'],
    'Department': ['Enterprise Sales', 'SMB Sales', 'Enterprise Sales', 'SMB Sales'],
    'Q1_Target': [50000, 30000, 55000, 25000],
    'Q1_Actual': [52000, 29000, 60000, 28000],
    'Performance': ['Exceeds', 'Needs Improvement', 'Outstanding', 'Exceeds']
}
df = pd.DataFrame(data)
df.to_excel("Sales_Data_2026.xlsx", index=False)

print("\n✅ 所有模拟测试文档生成完毕！你可以再次运行文档处理流水线了。")
