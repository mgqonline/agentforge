import os
import pandas as pd
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors

os.makedirs('document_processor/samples', exist_ok=True)

# 1. 生成海量 Excel 测试数据 (多 Sheet, 多行数据)
print("Generating Excel data...")
# Sheet 1: 员工绩效表 (50行数据)
emp_data = {
    '员工ID': [f'E{i:03d}' for i in range(1, 51)],
    '姓名': [f'测试员工_{i}' for i in range(1, 51)],
    '部门': ['技术部', '市场部', '人力资源部', '财务部', '销售部'] * 10,
    '绩效评分': [80.0 + (i % 20) for i in range(50)],
    '项目经验': [f'参与项目_Alpha_{i}' for i in range(50)]
}
df_emp = pd.DataFrame(emp_data)

# Sheet 2: 财务报表 (长文本单元格)
finance_data = {
    '季度': ['2023-Q1', '2023-Q2', '2023-Q3', '2023-Q4'],
    '营收(万)': [1500, 1800, 2100, 2500],
    '详细说明': [
        '第一季度主要受宏观经济复苏影响，产品线A销量大幅度上涨，导致整体营收超出预期约15%。其中华南区贡献了主要的增长份额。',
        '第二季度发布了新一代大模型平台，客户订阅数激增。但同时研发投入达到了顶峰，利润率有所压缩，但战略地位得到了巩固。',
        '第三季度进入传统淡季，我们调整了营销策略，优化了获客成本（CAC）。整体营收保持稳健增长，老客户复购率达到了惊人的 85%。',
        '第四季度开启了年终大促活动，软件与硬件结合的解决方案获得了政企客户的青睐，单笔超千万的大订单达到了3个。'
    ]
}
df_finance = pd.DataFrame(finance_data)

with pd.ExcelWriter('document_processor/samples/test_data_large.xlsx') as writer:
    df_emp.to_excel(writer, sheet_name='员工绩效', index=False)
    df_finance.to_excel(writer, sheet_name='季度财报', index=False)


# 2. 生成多页 Word 测试数据 (带复杂层级)
print("Generating Word data...")
doc = Document()
doc.add_heading('大型人工智能工程实践规范指南 (内部版)', 0)

doc.add_heading('第一章：引言与总体架构', level=1)
doc.add_paragraph('在现代软件开发中，大语言模型（LLM）的接入已经成为不可逆转的趋势。本指南旨在规范研发团队在接入 AI 时的标准流程。' * 5)
doc.add_paragraph('我们推崇的技术栈包括：Python 3.10+, FastAPI, LangChain, 向量数据库 ChromaDB 等。' * 5)

doc.add_heading('第二章：RAG (检索增强生成) 最佳实践', level=1)
doc.add_heading('2.1 文档解析与处理', level=2)
for i in range(3):
    doc.add_paragraph(f'这是第 {i+1} 条关于文档处理的规范：务必保证输入到大模型的数据是纯净的、没有噪音的。对于 PDF 文件的读取，我们推荐使用 PyMuPDF，因为它的文字定位更加准确。对于表格，我们需要将二维的结构序列化为 Markdown 表格，这样能够极大地帮助 LLM 理解行与列的对应关系，防止产生位置幻觉 (Spatial Hallucination)。' * 2)

doc.add_heading('2.2 向量检索与混合检索', level=2)
doc.add_paragraph('仅仅依赖向量检索 (Vector Search) 是不够的，因为向量模型往往缺乏对专有名词、序列号的精确匹配能力。因此，我们必须引入 BM25 作为词频基线的兜底方案。这种融合机制通常被称为 Hybrid Search。在取得双路召回结果后，还需要使用重排序模型 (CrossEncoder) 来提纯最终的候选结果。' * 3)

doc.save('document_processor/samples/test_report_large.docx')


# 3. 生成多页 PDF 测试数据
print("Generating PDF data...")
c = canvas.Canvas('document_processor/samples/test_paper_large.pdf', pagesize=letter)
width, height = letter

# 第一页
c.setFont("Helvetica-Bold", 16)
c.drawString(100, height - 100, "Advanced Retrieval-Augmented Generation (RAG)")
c.setFont("Helvetica", 12)
text_y = height - 140
for i in range(15):
    c.drawString(100, text_y, f"This is line {i+1} of the abstract. The LLM revolution has shown that providing relevant")
    text_y -= 20
    c.drawString(100, text_y, f"context dynamically allows models to overcome hallucination and temporal limitations.")
    text_y -= 25

c.showPage() # 换页

# 第二页
c.setFont("Helvetica-Bold", 14)
c.drawString(100, height - 100, "1. Architecture Overview")
c.setFont("Helvetica", 12)
text_y = height - 140
for i in range(20):
    c.drawString(100, text_y, f"[Architecture Detail {i+1}]: The data ingestion pipeline starts with a Unified Parser,")
    text_y -= 15
    c.drawString(100, text_y, f"which routes files to specific handlers based on their MIME types and extensions.")
    text_y -= 20

c.save()

print("Large sample files generated successfully.")
