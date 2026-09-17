import os
import json
import pandas as pd
from pptx import Presentation

samples_dir = "/Users/mac/Documents/project/ailearning/document_processor/samples"

# 1. TXT
with open(os.path.join(samples_dir, "test.txt"), "w", encoding="utf-8") as f:
    f.write("这是一个纯文本文件的测试用例。\n包含多行文本内容。\n主要用于测试基础解析。")

# 2. MD
with open(os.path.join(samples_dir, "test.md"), "w", encoding="utf-8") as f:
    f.write("# 大模型 RAG\n\n## 介绍\n这是关于 RAG 的一段测试 Markdown 文档。\n\n```python\nprint('hello')\n```")

# 3. CSV
df = pd.DataFrame({
    "姓名": ["张三", "李四", "王五"],
    "年龄": [25, 30, 28],
    "部门": ["AI 实验室", "研发部", "产品部"]
})
df.to_csv(os.path.join(samples_dir, "test.csv"), index=False)

# 4. JSON
data = {
    "project": "AI Terminal",
    "version": "1.0",
    "features": ["RAG", "Agent", "Multi-modal"]
}
with open(os.path.join(samples_dir, "test.json"), "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

# 5. HTML
html_content = """
<!DOCTYPE html>
<html>
<head>
<style> body { font-size: 14px; } </style>
<script> console.log("script block should be removed"); </script>
</head>
<body>
    <h1>欢迎使用 AI Terminal</h1>
    <p>这是一个 HTML 文档解析测试。它应该能自动去除脚本和样式表内容，仅提取网页里的纯文字。</p>
</body>
</html>
"""
with open(os.path.join(samples_dir, "test.html"), "w", encoding="utf-8") as f:
    f.write(html_content)

# 6. XML
xml_content = """<?xml version="1.0" encoding="UTF-8"?>
<note>
  <to>AI 助手</to>
  <from>用户</from>
  <heading>测试 XML 解析</heading>
  <body>请解析这里的节点内容！</body>
</note>
"""
with open(os.path.join(samples_dir, "test.xml"), "w", encoding="utf-8") as f:
    f.write(xml_content)

# 7. PPTX
prs = Presentation()
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "PPTX 解析测试标题"
subtitle.text = "这是第一页的副标题文本。验证 python-pptx 是否生效。"

slide_layout_2 = prs.slide_layouts[1]
slide_2 = prs.slides.add_slide(slide_layout_2)
title_2 = slide_2.shapes.title
body_2 = slide_2.placeholders[1]
title_2.text = "第二页：架构介绍"
body_2.text = "1. 前端\n2. 后端\n3. 数据库"

prs.save(os.path.join(samples_dir, "test.pptx"))

print("Test files generated successfully.")
